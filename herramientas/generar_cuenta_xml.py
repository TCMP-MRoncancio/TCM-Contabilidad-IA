#!/usr/bin/env python3
"""
Genera archivos XML de cuentas (chart of accounts) para Kondor K+/K+TP,
replicando EXACTAMENTE la salida de la macro `Accounts()` del archivo
`Cargador_Datos_v21.xlsm` (hoja "Account", modulo "Modulo2").

Uso:
    python herramientas/generar_cuenta_xml.py cuentas.json
    python herramientas/generar_cuenta_xml.py cuentas.json --forzar

El lote es atomico: se validan TODAS las cuentas primero (incluidos
duplicados dentro del mismo lote). Si una sola falla, no se escribe
ningun archivo. Un lote vacio tambien se rechaza (no hay nada que hacer).

Por defecto, si ya existe un .xml en /Account con el mismo
Account_ShortName, la cuenta se RECHAZA. Con --forzar se permite
sobrescribir, y se guarda una copia .bak con marca de tiempo del
archivo anterior antes de hacerlo (nunca pisa un respaldo previo).

Cada corrida exitosa actualiza tambien reportes/reporte_cuentas.xlsx,
consolidando TODAS las cuentas que existan en /Account en ese momento.
"""

import json
import shutil
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime
from xml.sax.saxutils import escape as escape_xml

try:
    import jsonschema
except ImportError:
    print("Falta la dependencia 'jsonschema'. Instalar con: python -m pip install -r requirements.txt")
    sys.exit(1)

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHEMA_PATH = REPO_ROOT / "schemas" / "cuenta.schema.json"
CATALOGO_PROYECTOS_PATH = REPO_ROOT / "schemas" / "catalogo_proyectos.json"
SALIDAS_DIR = REPO_ROOT / "Account"
REPORTES_DIR = REPO_ROOT / "reportes"
REPORTE_PATH = REPORTES_DIR / "reporte_cuentas.xlsx"

ACCOUNT_NAME_LARGO_MAXIMO = None

VALORES_CONFIRMADOS = {
    "AccountType": {"B"},
    "ValuationType": {"N"},
    "InputMode": {"C"},
}


def cargar_schema() -> dict:
    with open(SCHEMA_PATH, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def cargar_proyectos_validos() -> set:
    with open(CATALOGO_PROYECTOS_PATH, "r", encoding="utf-8-sig") as f:
        datos = json.load(f)
    return set(datos.get("proyectos_validos", []))


def validar_estructural(cuenta, schema: dict) -> list:
    """Valida tipos, patrones y additionalProperties usando jsonschema.
    Si esto falla, NO se corren las reglas de negocio."""
    validador = jsonschema.Draft7Validator(schema)
    errores = []
    for err in sorted(validador.iter_errors(cuenta), key=lambda e: list(e.path)):
        ruta = ".".join(str(p) for p in err.path) or "(raiz)"
        errores.append(f"{ruta}: {err.message}")
    return errores


def validar_reglas_negocio(cuenta: dict, permitir_duplicados: bool, nombres_vistos_en_lote: set) -> list:
    """Reglas que el JSON Schema no puede expresar."""
    errores = []

    nombre = cuenta.get("Account_Name", "")
    if nombre.strip() == "":
        errores.append("Account_Name no puede estar vacio ni contener solo espacios.")
    elif nombre != nombre.strip():
        errores.append(
            "Account_Name tiene espacios al inicio o al final ('" + nombre + "'). "
            "Esto esta pendiente de confirmar (¿la macro los recorta, rechaza o "
            "conserva?) - ver vault-patrones/candidatos/enums-pendientes.md. "
            "Quita los espacios sobrantes para continuar."
        )
    elif ACCOUNT_NAME_LARGO_MAXIMO is not None and len(nombre) > ACCOUNT_NAME_LARGO_MAXIMO:
        errores.append(
            f"Account_Name supera el largo maximo confirmado ({ACCOUNT_NAME_LARGO_MAXIMO} caracteres): "
            f"tiene {len(nombre)}."
        )

    for campo, valores_ok in VALORES_CONFIRMADOS.items():
        valor = cuenta.get(campo)
        if valor and valor not in valores_ok:
            errores.append(
                f"'{campo}' = '{valor}' no esta en la lista de valores confirmados "
                f"{sorted(valores_ok)}. Si es un valor nuevo y valido, agregalo a "
                f"vault-patrones/kondor/cuentas/estructura-cuenta.md y a VALORES_CONFIRMADOS "
                f"en este script antes de usarlo."
            )

    proyecto = cuenta.get("ChartOfAccount_Id")
    if proyecto:
        proyectos_validos = cargar_proyectos_validos()
        if proyecto not in proyectos_validos:
            errores.append(
                f"ChartOfAccount_Id = '{proyecto}' no esta en el catalogo de proyectos "
                f"confirmados {sorted(proyectos_validos)}. Agregalo primero a "
                f"vault-patrones/kondor/cuentas/catalogo-proyectos.md y "
                f"schemas/catalogo_proyectos.json via PR."
            )

    short_name = cuenta.get("Account_ShortName", "")
    if short_name:
        if short_name in nombres_vistos_en_lote:
            errores.append(
                f"Numero de cuenta '{short_name}' repetido dentro de este mismo lote "
                f"(ya aparece en una cuenta anterior de la misma solicitud)."
            )
        ruta_existente = SALIDAS_DIR / f"{short_name}.xml"
        if ruta_existente.exists() and not permitir_duplicados:
            errores.append(
                f"Ya existe una cuenta con el numero '{short_name}' en "
                f"{SALIDAS_DIR.name}/{ruta_existente.name}. No se genera para evitar "
                f"duplicados. Si de verdad quieres reemplazarla, vuelve a correr el "
                f"script agregando --forzar."
            )

    return errores


def generar_xml(cuenta: dict) -> str:
    """Todos los valores se escapan con xml.sax.saxutils.escape antes de
    insertarse, como defensa en profundidad."""
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        "<Account>\n"
        f'<Account_ShortName type="string">{escape_xml(str(cuenta["Account_ShortName"]))}</Account_ShortName>\n'
        f'<Account_Name type="string">{escape_xml(cuenta["Account_Name"])}</Account_Name>\n'
        f'<ChartOfAccount_Id type="string">{escape_xml(cuenta["ChartOfAccount_Id"])}</ChartOfAccount_Id>\n'
        f'<AccountType type="enum">{escape_xml(cuenta["AccountType"])}</AccountType>\n'
        f'<ValuationType type="enum">{escape_xml(cuenta["ValuationType"])}</ValuationType>\n'
        f'<InputMode type="enum">{escape_xml(cuenta["InputMode"])}</InputMode>\n'
        "</Account>\n"
    )


def _ruta_backup_sin_pisar(ruta: Path) -> Path:
    """Genera un nombre de respaldo con marca de tiempo (AAAAMMDD-HHMMSS).
    Si ya existe un respaldo con ese mismo segundo, agrega un sufijo
    numerico para no pisarlo nunca."""
    marca = datetime.now().strftime("%Y%m%d-%H%M%S")
    candidato = ruta.with_suffix(ruta.suffix + f".{marca}.bak")
    contador = 2
    while candidato.exists():
        candidato = ruta.with_suffix(ruta.suffix + f".{marca}-{contador}.bak")
        contador += 1
    return candidato


def escribir_cuenta(cuenta: dict, permitir_duplicados: bool) -> Path:
    """Escribe el .xml con newline='\\r\\n' explicito (confirmado contra
    11 casos reales exportados por la macro, ver Bloque 3). Si ya existe
    y se permite sobrescribir, guarda una copia .bak con marca de tiempo
    antes, sin pisar respaldos anteriores."""
    SALIDAS_DIR.mkdir(exist_ok=True)
    ruta = SALIDAS_DIR / f"{cuenta['Account_ShortName']}.xml"

    if ruta.exists() and permitir_duplicados:
        ruta_bak = _ruta_backup_sin_pisar(ruta)
        shutil.copy2(ruta, ruta_bak)
        print(f"   (copia de respaldo: {ruta_bak.name})")

    ruta.write_text(generar_xml(cuenta), encoding="utf-8", newline="\r\n")
    return ruta


def actualizar_reporte():
    """Escanea TODOS los .xml en /Account y regenera el Excel consolidado."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("Aviso: openpyxl no esta instalado, se omite el reporte Excel.")
        print("Instalalo con: python -m pip install openpyxl")
        return

    if not SALIDAS_DIR.exists():
        return

    columnas = [
        "Account_ShortName", "Account_Name", "ChartOfAccount_Id",
        "AccountType", "ValuationType", "InputMode", "Archivo",
    ]

    filas = []
    for xml_file in sorted(SALIDAS_DIR.glob("*.xml")):
        try:
            root = ET.parse(xml_file).getroot()
            fila = [
                (root.findtext("Account_ShortName") or ""),
                (root.findtext("Account_Name") or ""),
                (root.findtext("ChartOfAccount_Id") or ""),
                (root.findtext("AccountType") or ""),
                (root.findtext("ValuationType") or ""),
                (root.findtext("InputMode") or ""),
                xml_file.name,
            ]
            filas.append(fila)
        except ET.ParseError:
            print(f"Aviso: no se pudo leer {xml_file.name} para el reporte (XML invalido).")

    REPORTES_DIR.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuentas"

    ws.append(columnas)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="16233F", end_color="16233F", fill_type="solid")

    for fila in filas:
        ws.append(fila)

    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    ws_meta = wb.create_sheet("Info")
    ws_meta.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws_meta.append(["Total de cuentas", len(filas)])
    ws_meta.append(["Fuente", "reportes/reporte_cuentas.xlsx generado automaticamente desde /Account. No editar a mano; se sobrescribe en cada corrida."])

    wb.save(REPORTE_PATH)
    print(f"OK Reporte actualizado: {REPORTE_PATH} ({len(filas)} cuentas)")


def main():
    args = [a for a in sys.argv[1:] if a != "--forzar"]
    permitir_duplicados = "--forzar" in sys.argv

    if len(args) != 1:
        print("Uso: python herramientas/generar_cuenta_xml.py cuentas.json [--forzar]")
        sys.exit(1)

    entrada = Path(args[0])
    if not entrada.exists():
        print(f"No se encontro el archivo: {entrada}")
        sys.exit(1)

    with open(entrada, "r", encoding="utf-8-sig") as f:
        datos = json.load(f)

    cuentas = datos if isinstance(datos, list) else [datos]

    if len(cuentas) == 0:
        print("XX El lote esta vacio: no hay cuentas para procesar.")
        sys.exit(1)

    schema = cargar_schema()

    errores_por_cuenta = {}
    nombres_vistos_en_lote = set()

    for i, cuenta in enumerate(cuentas):
        if not isinstance(cuenta, dict):
            errores_por_cuenta[i] = [
                f"posicion {i}: se esperaba un objeto (diccionario) con los campos de la "
                f"cuenta, se recibio {type(cuenta).__name__}."
            ]
            continue

        errores = validar_estructural(cuenta, schema)
        if not errores:
            errores = validar_reglas_negocio(cuenta, permitir_duplicados, nombres_vistos_en_lote)

        if errores:
            errores_por_cuenta[i] = errores
        else:
            nombres_vistos_en_lote.add(cuenta["Account_ShortName"])

    if errores_por_cuenta:
        print(f"XX Lote RECHAZADO: {len(errores_por_cuenta)}/{len(cuentas)} cuenta(s) con errores.")
        print("   No se genero ningun archivo (el lote es atomico: todo o nada).\n")
        for i, errores in errores_por_cuenta.items():
            cuenta_i = cuentas[i]
            identificador = cuenta_i.get("Account_ShortName", f"posicion {i}") if isinstance(cuenta_i, dict) else f"posicion {i}"
            print(f"Cuenta '{identificador}':")
            for e in errores:
                print(f"   - {e}")
        sys.exit(1)

    for cuenta in cuentas:
        ruta = escribir_cuenta(cuenta, permitir_duplicados)
        print(f"OK Generado: {ruta}")

    print(f"\n{len(cuentas)}/{len(cuentas)} cuentas generadas correctamente.")

    actualizar_reporte()
    sys.exit(0)


if __name__ == "__main__":
    main()