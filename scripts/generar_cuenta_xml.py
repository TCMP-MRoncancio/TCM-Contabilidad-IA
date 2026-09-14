#!/usr/bin/env python3
"""
Genera archivos XML de cuentas (chart of accounts) para Kondor K+/K+TP,
replicando EXACTAMENTE la salida de la macro `Accounts()` del archivo
`Cargador_Datos_v21.xlsm` (hoja "Account", modulo "Modulo2").

Uso:
    python scripts/generar_cuenta_xml.py cuentas.json

Cada cuenta genera un archivo <Account_ShortName>.xml en /Account.
Al final de cada corrida, tambien actualiza reportes/reporte_cuentas.xlsx
consolidando TODAS las cuentas que existan en /Account en ese momento
(no solo las de esta corrida) - asi el reporte siempre queda sincronizado
con lo que realmente hay en la carpeta, sin importar si se genero 1 o 1000
cuentas, en una corrida o en varias.
"""

import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHEMA_PATH = REPO_ROOT / "schemas" / "cuenta.schema.json"
CATALOGO_PROYECTOS_PATH = REPO_ROOT / "schemas" / "catalogo_proyectos.json"
SALIDAS_DIR = REPO_ROOT / "Account"
REPORTES_DIR = REPO_ROOT / "reportes"
REPORTE_PATH = REPORTES_DIR / "reporte_cuentas.xlsx"

# Largo fijo esperado para Account_ShortName - todas las cuentas del catalogo deben coincidir
LARGO_ESPERADO_SHORT_NAME = 16

# Valores confirmados para los enums de estructura fija (no varian por proyecto).
VALORES_CONFIRMADOS = {
    "AccountType": {"B"},
    "ValuationType": {"N"},
    "InputMode": {"C"},
}


def cargar_schema():
    with open(SCHEMA_PATH, "r", encoding="utf-8-sig") as f:
        return json.load(f)


def cargar_proyectos_validos() -> set:
    with open(CATALOGO_PROYECTOS_PATH, "r", encoding="utf-8-sig") as f:
        datos = json.load(f)
    return set(datos.get("proyectos_validos", []))


def validar(cuenta: dict, schema: dict) -> list:
    errores = []

    for campo in schema["required"]:
        if campo not in cuenta or cuenta[campo] in (None, ""):
            errores.append(f"Falta el campo obligatorio: {campo}")

    short_name = cuenta.get("Account_ShortName", "")
    if short_name and not (short_name.isdigit() and len(short_name) == LARGO_ESPERADO_SHORT_NAME):
        errores.append(
            f"Account_ShortName debe ser numerico de {LARGO_ESPERADO_SHORT_NAME} digitos "
            f"(mismo largo que las demas cuentas del catalogo), se recibio: "
            f"'{short_name}' ({len(short_name)} caracteres)"
        )

    nombre = cuenta.get("Account_Name", "")
    if nombre and not re.match(r"^[A-Za-z0-9 ]+$", nombre):
        errores.append(
            f"Account_Name '{nombre}' contiene tildes o caracteres especiales no permitidos. "
            f"Solo se permiten letras sin acento, numeros y espacios."
        )

    for campo, valores_ok in VALORES_CONFIRMADOS.items():
        valor = cuenta.get(campo)
        if valor and valor not in valores_ok:
            errores.append(
                f"'{campo}' = '{valor}' no esta en la lista de valores confirmados "
                f"{sorted(valores_ok)}. Si es un valor nuevo y valido, agregalo a "
                f"docs/estructura-cuenta.md y a VALORES_CONFIRMADOS en este script "
                f"antes de usarlo."
            )

    proyecto = cuenta.get("ChartOfAccount_Id")
    if proyecto:
        proyectos_validos = cargar_proyectos_validos()
        if proyecto not in proyectos_validos:
            errores.append(
                f"ChartOfAccount_Id = '{proyecto}' no esta en el catalogo de proyectos "
                f"confirmados {sorted(proyectos_validos)}. Agregalo primero a "
                f"docs/catalogo-proyectos.md y schemas/catalogo_proyectos.json via PR."
            )

    return errores


def generar_xml(cuenta: dict) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        "<Account>\n"
        f'<Account_ShortName type="string">{cuenta["Account_ShortName"]}</Account_ShortName>\n'
        f'<Account_Name type="string">{cuenta["Account_Name"]}</Account_Name>\n'
        f'<ChartOfAccount_Id type="string">{cuenta["ChartOfAccount_Id"]}</ChartOfAccount_Id>\n'
        f'<AccountType type="enum">{cuenta["AccountType"]}</AccountType>\n'
        f'<ValuationType type="enum">{cuenta["ValuationType"]}</ValuationType>\n'
        f'<InputMode type="enum">{cuenta["InputMode"]}</InputMode>\n'
        "</Account>\n"
    )


def procesar_cuenta(cuenta: dict, schema: dict) -> bool:
    errores = validar(cuenta, schema)
    if errores:
        print(f"XX Cuenta '{cuenta.get('Account_ShortName', '?')}' NO paso la validacion:")
        for e in errores:
            print(f"   - {e}")
        return False

    SALIDAS_DIR.mkdir(exist_ok=True)
    ruta = SALIDAS_DIR / f"{cuenta['Account_ShortName']}.xml"
    ruta.write_text(generar_xml(cuenta), encoding="utf-8")
    print(f"OK Generado: {ruta}")
    return True


def actualizar_reporte():
    """Escanea TODOS los .xml en /Account y regenera el Excel consolidado.
    No modifica los XML - solo los lee para armar la tabla resumen."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("Aviso: openpyxl no esta instalado, se omite el reporte Excel.")
        print("Instalalo con: python -m pip install openpyxl")
        return

    if not SALIDAS_DIR.exists():
        return

    columnas = [
        "Account_ShortName", "Account_Name", "ChartOfAccount_Id",
        "AccountType", "ValuationType", "InputMode", "Archivo",
    ]

    filas = []
    for xml_file in sorted(SALIDAS_DIR.glob("*.xml")):
        try:
            root = ET.parse(xml_file).getroot()
            fila = [
                (root.findtext("Account_ShortName") or ""),
                (root.findtext("Account_Name") or ""),
                (root.findtext("ChartOfAccount_Id") or ""),
                (root.findtext("AccountType") or ""),
                (root.findtext("ValuationType") or ""),
                (root.findtext("InputMode") or ""),
                xml_file.name,
            ]
            filas.append(fila)
        except ET.ParseError:
            print(f"Aviso: no se pudo leer {xml_file.name} para el reporte (XML invalido).")

    REPORTES_DIR.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuentas"

    ws.append(columnas)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="16233F", end_color="16233F", fill_type="solid")

    for fila in filas:
        ws.append(fila)

    for col in ws.columns:
        max_len = max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    ws_meta = wb.create_sheet("Info")
    ws_meta.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws_meta.append(["Total de cuentas", len(filas)])
    ws_meta.append(["Fuente", "reportes/reporte_cuentas.xlsx generado automaticamente desde /Account. No editar a mano; se sobrescribe en cada corrida."])

    wb.save(REPORTE_PATH)
    print(f"OK Reporte actualizado: {REPORTE_PATH} ({len(filas)} cuentas)")


def main():
    if len(sys.argv) != 2:
        print("Uso: python scripts/generar_cuenta_xml.py cuentas.json")
        sys.exit(1)

    entrada = Path(sys.argv[1])
    if not entrada.exists():
        print(f"No se encontro el archivo: {entrada}")
        sys.exit(1)

    with open(entrada, "r", encoding="utf-8-sig") as f:
        datos = json.load(f)

    cuentas = datos if isinstance(datos, list) else [datos]
    schema = cargar_schema()

    ok = sum(procesar_cuenta(c, schema) for c in cuentas)
    print(f"\n{ok}/{len(cuentas)} cuentas generadas correctamente.")

    actualizar_reporte()

    sys.exit(0 if ok == len(cuentas) else 1)


if __name__ == "__main__":
    main()